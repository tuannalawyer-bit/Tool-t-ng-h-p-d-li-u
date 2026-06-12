import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_base_template(output_path, theme="blank"):
    prs = Presentation()
    
    # Set presentation dimensions to 16:9 widescreen (standard for modern projectors)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout to build our custom framework
    blank_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(blank_layout)
    
    # Masan Brand Color Hex Definitions
    MASAN_RED = RGBColor(220, 30, 40)   # Corporate red
    MASAN_GOLD = RGBColor(210, 170, 80)  # Elegance Gold Accent
    PURE_WHITE = RGBColor(255, 255, 255)
    DARK_TEXT = RGBColor(40, 40, 45)
    LIGHT_GRAY = RGBColor(240, 240, 242)

    # 1. Draw Theme Background / Elements
    if theme == "masan":
        # Add Top Header Bar
        header_rect = slide.shapes.add_shape(
            1, # MSO_SHAPE.RECTANGLE
            Inches(0), Inches(0), prs.slide_width, Inches(1.2)
        )
        header_rect.fill.solid()
        header_rect.fill.fore_color.rgb = MASAN_RED
        header_rect.line.fill.background() # Remove border
        
        # Add thin gold separator strip
        sep_rect = slide.shapes.add_shape(
            1, Inches(0), Inches(1.2), prs.slide_width, Inches(0.1)
        )
        sep_rect.fill.solid()
        sep_rect.fill.fore_color.rgb = MASAN_GOLD
        sep_rect.line.fill.background()
        
        # Setup Title coordinates to be inside header
        title_left, title_top = Inches(0.5), Inches(0.2)
        title_width, title_height = Inches(12.33), Inches(0.8)
        title_color = PURE_WHITE
        
        # Setup Content coordinates below header
        content_left, content_top = Inches(0.8), Inches(1.8)
        content_width, content_height = Inches(11.73), Inches(4.8)
        
        # Add Footer decoration
        footer_rect = slide.shapes.add_shape(
            1, Inches(0.5), Inches(7.0), Inches(3.0), Inches(0.05)
        )
        footer_rect.fill.solid()
        footer_rect.fill.fore_color.rgb = MASAN_GOLD
        footer_rect.line.fill.background()
        
    else: # theme == "blank"
        # Setup standard corporate layout coordinates
        title_left, title_top = Inches(0.8), Inches(0.5)
        title_width, title_height = Inches(11.73), Inches(1.0)
        title_color = DARK_TEXT
        
        content_left, content_top = Inches(0.8), Inches(1.8)
        content_width, content_height = Inches(11.73), Inches(4.8)

    # 2. Add Title Textbox
    title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = "{{TIEU_DE}}"
    p_title.font.bold = True
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(40)
    p_title.font.color.rgb = title_color
    p_title.alignment = PP_ALIGN.LEFT

    # 3. Add Body Content Textbox
    content_box = slide.shapes.add_textbox(content_left, content_top, content_width, content_height)
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    p_content = tf_content.paragraphs[0]
    p_content.text = "{{NOI_DUNG}}"
    p_content.font.name = "Calibri"
    p_content.font.size = Pt(20)
    p_content.font.color.rgb = DARK_TEXT
    
    # Save result
    prs.save(output_path)
    print(f"Generated template '{theme}' at: {output_path}")

def ensure_all_templates(target_dir="templates"):
    if not os.path.exists(target_dir):
        os.makedirs(target_dir)
        
    blank_path = os.path.join(target_dir, "Default_Blank.pptx")
    masan_path = os.path.join(target_dir, "Masan_Brand_Template.pptx")
    
    if not os.path.exists(blank_path):
        create_base_template(blank_path, "blank")
    if not os.path.exists(masan_path):
        create_base_template(masan_path, "masan")

if __name__ == "__main__":
    ensure_all_templates()
