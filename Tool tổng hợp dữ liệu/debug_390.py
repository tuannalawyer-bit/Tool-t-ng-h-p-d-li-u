import openpyxl
import re

def debug_390():
    wb = openpyxl.load_workbook("Ket_Qua_Phan_Tich_Rui_Ro_Ton_Kho.xlsx")
    ws1 = wb["Top UU TIEN KIEM TRA"]
    
    target_row = None
    for row in ws1.iter_rows(min_row=4, values_only=True):
        if row[3] == 390 or row[1] == 'N/A' or 'N/A' in str(row[2]):
            print(f"Found matching row: Rank={row[0]}, ID={row[1]}, Name={row[2]}, Score={row[3]}, SKUs={row[4]}, StockVal={row[5]}, CLKK={row[6]}, STOCount={row[7]}, STOVal={row[8]}")
            if row[3] == 390:
                target_row = row

    print("\nLet's inspect 'Chi Tiet Lech Kiem Ke' sheet for any entries mapped to N/A or similar:")
    ws4 = wb["Chi Tiet Lech Kiem Ke"]
    count_na = 0
    for row in ws4.iter_rows(min_row=2, values_only=True):
        if row[0] == 'N/A' or row[0] == 'WM+' or len(str(row[0])) > 8: # Pseudo IDs might have len > 8
            count_na += 1
            if count_na <= 5:
                print(f"CLKK row: ID={row[0]}, Location={row[1]}, Product={row[2]}, Warning={row[4]}")
    print(f"Total unmatched CLKK rows: {count_na}")

if __name__ == "__main__":
    debug_390()
