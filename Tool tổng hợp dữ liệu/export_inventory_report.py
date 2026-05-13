import zipfile
import xml.etree.ElementTree as ET
import sys
import os
import re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def load_shared_strings(z):
    shared_strings = []
    try:
        with z.open('xl/sharedStrings.xml') as f:
            for event, elem in ET.iterparse(f, events=('end',)):
                if elem.tag.endswith('}si'):
                    text = "".join(elem.itertext())
                    shared_strings.append(text)
                    elem.clear()
    except KeyError:
        pass
    return shared_strings

def get_cell_value(elem, cell_type, shared_strings):
    val_elem = elem.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v')
    if val_elem is None or val_elem.text is None:
        return None
    val = val_elem.text
    if cell_type == 's':
        try:
            idx = int(val)
            return shared_strings[idx] if idx < len(shared_strings) else f"SS_{idx}"
        except:
            return val
    return val

def parse_sheet_all_rows(z, xml_path, shared_strings):
    rows = []
    try:
        with z.open('xl/' + xml_path) as f:
            context = ET.iterparse(f, events=('start', 'end'))
            current_row_data = {}
            
            for event, elem in context:
                tag = elem.tag
                if event == 'start':
                    if tag.endswith('}row'):
                        current_row_data = {}
                elif event == 'end':
                    if tag.endswith('}c'):
                        cell_ref = elem.attrib.get('r')
                        cell_type = elem.attrib.get('t')
                        val = get_cell_value(elem, cell_type, shared_strings)
                        if val is not None and cell_ref:
                            col = re.sub(r'\d+', '', cell_ref)
                            current_row_data[col] = val
                    elif tag.endswith('}row'):
                        if current_row_data:
                            rows.append(current_row_data)
                        elem.clear()
    except Exception as e:
        print(f"Error parsing {xml_path}: {e}")
    return rows

def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except:
        pass

    filepath = "Layer_1 4.xlsx"
    if not os.path.exists(filepath):
        print(f"Source file not found: {filepath}")
        return

    print("=== BUILING INVENTORY RISK ANALYSIS & EXCEL EXPORT SYSTEM ===")
    
    with zipfile.ZipFile(filepath, 'r') as z:
        # Get Sheet Map
        rel_map = {}
        with z.open('xl/_rels/workbook.xml.rels') as f:
            tree = ET.parse(f)
            root = tree.getroot()
            ns = {'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
            for rel in root.findall('.//r:Relationship', ns):
                rel_map[rel.attrib.get('Id')] = rel.attrib.get('Target')

        sheets = {}
        with z.open('xl/workbook.xml') as f:
            tree = ET.parse(f)
            root = tree.getroot()
            ns = {'ns': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            for sheet in root.findall('.//ns:sheet', ns):
                sname = sheet.attrib.get('name')
                rid = sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                xml_path = rel_map.get(rid, "Unknown")
                sheets[sname] = xml_path

        shared_strings = load_shared_strings(z)

        # --- STEP 1: BUILD STORE MASTER MAPPING ---
        print("Building Store ID-to-Name map from List_Store_V2 and STO sheets...")
        store_map = {} # ID -> Name
        
        # Map from STO Treo sheet (accurate columns)
        if 'STO_Treo_>10tr' in sheets:
            sto_rows = parse_sheet_all_rows(z, sheets['STO_Treo_>10tr'], shared_strings)
            for row in sto_rows:
                sid = row.get('A')
                sname = row.get('B')
                if sid and sname and sname.startswith('WM+'):
                    store_map[str(sid).strip()] = str(sname).strip()

        # Scan List_Store_V2 for supplemental mappings
        if 'List_Store_V2' in sheets:
            ls_rows = parse_sheet_all_rows(z, sheets['List_Store_V2'], shared_strings)
            for row in ls_rows:
                # Find numeric-like keys (4 digits or 2XXX/GXXX)
                potential_ids = []
                potential_names = []
                for val in row.values():
                    sval = str(val).strip()
                    if re.match(r'^\d{4}$', sval) or re.match(r'^[2G][A-Za-z0-9]{3}$', sval):
                        potential_ids.append(sval)
                    elif sval.startswith('WM+') or sval.startswith('WIN') or sval.startswith('WinMart'):
                        potential_names.append(sval)
                
                # Map closest matches if possible
                # For simple mapping, we store them in store_map
                for sid in potential_ids:
                    for sname in potential_names:
                        # If a store ID isn't mapped yet, or we can infer
                        if sid not in store_map and len(potential_ids) == 1 and len(potential_names) == 1:
                            store_map[sid] = sname
        
        # Helper to search store name from a string or fallback to dictionary
        def normalize_store(input_str):
            # Extract ID using regex if input is a long description
            id_match = re.search(r'tại CH\s+([A-Za-z0-9]+)', input_str)
            if id_match:
                sid = id_match.group(1)
                return sid, store_map.get(sid, f"Store ID {sid}")
            
            # Check if input_str itself is an ID in map
            cleaned = str(input_str).strip()
            if cleaned in store_map:
                return cleaned, store_map[cleaned]
            
            # Check if input_str contains store ID
            for sid, name in store_map.items():
                if sid in cleaned:
                    return sid, name
                    
            # If it is a store name but not in map, try to match
            if cleaned.startswith('WM+') or cleaned.startswith('WIN'):
                # Maybe reverse lookup? Or generate a pseudo ID
                for sid, name in store_map.items():
                    if cleaned in name or name in cleaned:
                        return sid, name
                return "N/A", cleaned
            
            return cleaned, "N/A"

        # --- STEP 2: AGGREGATE INVENTORY METRICS ---
        print("\nExtracting detailed risk data...")
        
        # Data storage structures
        store_scores = defaultdict(lambda: {
            'id': '', 'name': '',
            'score': 0,
            'ton_cao_count': 0, 'ton_cao_value': 0.0,
            'clkk_count': 0,
            'sto_treo_count': 0, 'sto_treo_value': 0.0,
            'sto_lech_count': 0
        })
        
        detail_ton_cao = []
        detail_sto_treo = []
        detail_clkk = []
        
        # A. Parse Ton_cao
        print("Parsing Ton_cao...")
        tc_rows = parse_sheet_all_rows(z, sheets['Ton_cao'], shared_strings)
        for row in tc_rows:
            desc = ""
            prod = ""
            for val in row.values():
                sval = str(val)
                if 'tồn cuối =' in sval:
                    desc = sval
                elif any(kw in sval for kw in ['NAM NGƯ', 'CHIN-SU', 'OMACHI', 'CUSTAS', 'Bia', 'Dầu']) or len(sval) > 15 and ' ' in sval and not 'tại CH' in sval:
                    # Guess product name based on content
                    if len(sval) < 50:
                        prod = sval
            
            if desc:
                sid_match = re.search(r'tại CH\s+([A-Za-z0-9]+)', desc)
                val_match = re.search(r'tồn cuối\s*=\s*(\d+)', desc)
                days_match = re.search(r'chưa bán\s*=\s*(\d+)', desc)
                
                sid = sid_match.group(1) if sid_match else "Unknown"
                val_amt = float(val_match.group(1)) if val_match else 0.0
                days = int(days_match.group(1)) if days_match else 0
                
                sname = store_map.get(sid, f"Store {sid}")
                
                detail_ton_cao.append({
                    'Store ID': sid,
                    'Store Name': sname,
                    'Product Name': prod if prod else "Top 50 SKU",
                    'Value (VND)': val_amt,
                    'Days Unsold': days,
                    'Risk Level': 'High' if days >= 30 else 'Medium',
                    'Description': desc
                })
                
                # Score Calculation Logic: 8 pts base, +2 pts if >=30 days
                pts = 8
                if days >= 30:
                    pts += 2
                    
                store_scores[sid]['id'] = sid
                store_scores[sid]['name'] = sname
                store_scores[sid]['ton_cao_count'] += 1
                store_scores[sid]['ton_cao_value'] += val_amt
                store_scores[sid]['score'] += pts

        # B. Parse STO_Treo_>10tr
        print("Parsing STO_Treo_>10tr...")
        sto_rows = parse_sheet_all_rows(z, sheets['STO_Treo_>10tr'], shared_strings)
        for row in sto_rows:
            sid = str(row.get('A', '')).strip()
            sname = str(row.get('B', '')).strip()
            val_amt = 0.0
            for v in row.values():
                try:
                    num = float(v)
                    if 10000000 <= num < 10000000000:
                        val_amt = max(val_amt, num)
                except: pass
                
            if sid and val_amt > 0:
                if not sname or sname == '':
                    sname = store_map.get(sid, f"Store {sid}")
                else:
                    store_map[sid] = sname # cache it
                
                # Phân nhóm rủi ro
                level = "Cần xử lý"
                if val_amt >= 100000000:
                    level = "Rất cao"
                elif val_amt >= 50000000:
                    level = "Cao"
                
                detail_sto_treo.append({
                    'Store ID': sid,
                    'Store Name': sname,
                    'Pending Value (VND)': val_amt,
                    'Risk Level': level
                })
                
                # Scoring: 9 pts base + scale bonuses
                pts = 9
                if val_amt >= 100000000: # > 100M
                    pts += 3
                if val_amt >= 1000000000: # > 1B
                    pts += 10
                    
                store_scores[sid]['id'] = sid
                store_scores[sid]['name'] = sname
                store_scores[sid]['sto_treo_count'] += 1
                store_scores[sid]['sto_treo_value'] += val_amt
                store_scores[sid]['score'] += pts

        # C. Parse CLKK_20%
        print("Parsing CLKK_20%...")
        clkk_rows = parse_sheet_all_rows(z, sheets['CLKK_20%'], shared_strings)
        for row in clkk_rows:
            # Skip header row
            if row.get('B') == 'store_id':
                continue
                
            sid = str(row.get('B', '')).strip()
            sname = str(row.get('C', '')).strip()
            prod = str(row.get('M', '')).strip()
            desc = str(row.get('AC', '')).strip()
            
            # Get estimated adjustment value from numerical columns if possible
            val_amt = 0.0
            for col in ['N', 'O', 'P', 'Q']:
                v = row.get(col)
                if v:
                    try:
                        num = abs(float(v))
                        if num > 0:
                            val_amt = max(val_amt, num)
                    except: pass

            if sid and sid != '' and desc and 'Cảnh báo' in desc:
                # Map to our score dictionary
                if sid not in store_map and sname:
                    store_map[sid] = sname

                detail_clkk.append({
                    'Store ID': sid,
                    'Store Location Name': sname,
                    'Product': prod,
                    'Est Adj Value (VND)': val_amt,
                    'Warning Type': desc
                })
                
                # Scoring: 10 pts per incident
                store_scores[sid]['id'] = sid
                if sname:
                    store_scores[sid]['name'] = sname
                store_scores[sid]['clkk_count'] += 1
                store_scores[sid]['score'] += 10

        # D. Parse STO_LechChuoi
        if 'STO_LechChuoi' in sheets:
            print("Parsing STO_LechChuoi...")
            lc_rows = parse_sheet_all_rows(z, sheets['STO_LechChuoi'], shared_strings)
            for row in lc_rows:
                sid = row.get('A', 'N/A')
                if sid != 'N/A':
                    store_scores[sid]['sto_lech_count'] += 1
                    store_scores[sid]['score'] += 9

        # --- STEP 3: BUILD EXCEL WORKBOOK ---
        print("\nBuilding Excel Report File using OpenPyXL...")
        wb = openpyxl.Workbook()
        
        # Create Styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, name="Arial", size=11)
        bold_font = Font(bold=True, name="Arial")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        currency_fmt = '#,##0'
        
        # -- SHEET 1: PRIORITY LIST (DANH SACH UU TIEN) --
        ws1 = wb.active
        ws1.title = "Top UU TIEN KIEM TRA"
        ws1.append(["XẾP HẠNG NGUY CƠ TỒN KHO - DANH SÁCH ƯU TIÊN THANH TRA"])
        ws1.merge_cells("A1:I1")
        ws1["A1"].font = Font(name="Arial", size=16, bold=True, color="C00000")
        ws1.append([])
        
        headers1 = [
            "Hạng", "Mã Cửa Hàng", "Tên Cửa Hàng", 
            "TỔNG ĐIỂM RỦI RO", "Số SKU Tồn Cao", "Giá Trị Tồn Đọng", 
            "Số Lần Lệch Kiểm Kê", "Lệnh STO Treo", "Giá Trị STO Treo"
        ]
        ws1.append(headers1)
        
        # Format Headers
        for cell in ws1[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        # Write ranked stores
        sorted_stores = sorted(store_scores.items(), key=lambda x: x[1]['score'], reverse=True)
        rank = 1
        for sid, d in sorted_stores:
            if d['score'] == 0: continue
            
            # Fill in Name if missing
            name = d['name']
            if not name or name == 'N/A' or name.startswith('Store '):
                name = store_map.get(sid, name)

            row_data = [
                rank, sid, name, 
                d['score'], d['ton_cao_count'], d['ton_cao_value'],
                d['clkk_count'], d['sto_treo_count'], d['sto_treo_value']
            ]
            ws1.append(row_data)
            rank += 1
            
        # Apply data styling
        for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row):
            row[0].alignment = center_align
            row[1].alignment = center_align
            row[3].font = bold_font
            row[3].alignment = center_align
            row[5].number_format = currency_fmt
            row[8].number_format = currency_fmt
            for cell in row:
                cell.border = thin_border

        # Highlight highest risk rows (top 10 in light red)
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        for r_idx in range(4, min(ws1.max_row + 1, 14)):
            for cell in ws1[r_idx]:
                cell.fill = red_fill

        # -- SHEET 2: DETAIL TON CAO --
        ws2 = wb.create_sheet(title="Chi Tiet Ton Cao")
        headers2 = ["Mã CH", "Tên CH", "Tên Sản Phẩm", "Giá Trị Tồn (VND)", "Số Ngày Chưa Bán", "Mức Độ Rủi Ro", "Mô Tả Chi Tiết"]
        ws2.append(headers2)
        for cell in ws2[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align; cell.border = thin_border
            
        for d in detail_ton_cao:
            ws2.append([d['Store ID'], d['Store Name'], d['Product Name'], d['Value (VND)'], d['Days Unsold'], d['Risk Level'], d['Description']])
            
        for row in ws2.iter_rows(min_row=2):
            row[3].number_format = currency_fmt
            row[4].alignment = center_align
            row[5].alignment = center_align
            for cell in row: cell.border = thin_border

        # -- SHEET 3: DETAIL STO TREO --
        ws3 = wb.create_sheet(title="Chi Tiet STO Treo")
        headers3 = ["Mã CH", "Tên CH", "Giá Trị Lệnh Treo (VND)", "Phân Nhóm Cảnh Báo"]
        ws3.append(headers3)
        for cell in ws3[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align; cell.border = thin_border
            
        for d in detail_sto_treo:
            ws3.append([d['Store ID'], d['Store Name'], d['Pending Value (VND)'], d['Risk Level']])
            
        for row in ws3.iter_rows(min_row=2):
            row[2].number_format = currency_fmt
            row[3].alignment = center_align
            for cell in row: cell.border = thin_border

        # -- SHEET 4: DETAIL CLKK --
        ws4 = wb.create_sheet(title="Chi Tiet Lech Kiem Ke")
        headers4 = ["Mã CH", "Tên Vị Trí / Cửa Hàng", "Sản Phẩm", "Ước Tính Giá Trị Lệch", "Loại Cảnh Báo Biến Động"]
        ws4.append(headers4)
        for cell in ws4[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align; cell.border = thin_border
            
        for d in detail_clkk:
            ws4.append([d['Store ID'], d['Store Location Name'], d['Product'], d['Est Adj Value (VND)'], d['Warning Type']])
            
        for row in ws4.iter_rows(min_row=2):
            row[3].number_format = currency_fmt
            for cell in row: cell.border = thin_border

        # Autofit columns
        for ws in [ws1, ws2, ws3, ws4]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                # set width, cap at 50
                ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

        # Save file
        out_file = "Ket_Qua_Phan_Tich_Rui_Ro_Ton_Kho.xlsx"
        wb.save(out_file)
        print(f"\nSUCCESFULLY CREATED REPORT FILE: {out_file}")
        print("Total Stores Evaluated:", len(store_scores))
        print(f"Output location: {os.path.abspath(out_file)}")

if __name__ == "__main__":
    main()
