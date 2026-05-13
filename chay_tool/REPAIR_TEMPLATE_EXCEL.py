import openpyxl
import os

paths = [
    r"D:\chạy tool\chạy tool\form tool.xlsx",
    r"d:\Dự án AI\chay_tool\form tool.xlsx"
]

for p in paths:
    if os.path.exists(p):
        try:
            wb = openpyxl.load_workbook(p)
            sheet_names = ["01. Thẩm định giá & HSPL"]
            for sn in sheet_names:
                if sn in wb.sheetnames:
                    ws = wb[sn]
                    # Check column 13 (M) header before deletion to confirm
                    header_val = ws.cell(row=3, column=13).value # Row 3 might be the header row
                    print(f"Checking {p} - col 13 header: {header_val}")
                    
                    # Actually delete Column 13
                    ws.delete_cols(13)
                    print(f"SUCCESS: Deleted column 13 from sheet {sn} in {p}")
            wb.save(p)
        except Exception as e:
            print(f"FAILED to process {p}: {e}")
