import openpyxl
import os
import unicodedata
import re
import sys
import io
import shutil

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

# --- CONFIG ---
FILE_CONTRACT = r"d:\Dự án AI\chay_tool\scratch\temp_contract.xlsx"
FILE_APPRAISAL = r"d:\Dự án AI\chay_tool\scratch\temp_appraisal.xlsx"
OUTPUT_FILE = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\Ket_Qua_Gop_Du_Lieu_V2.xlsx"

def remove_accents(input_str):
    if not input_str: return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    s = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    return s.replace('đ', 'd').replace('Đ', 'D')

def get_core_keywords(text):
    if not text: return set()
    text = remove_accents(text).lower()
    # Remove noise words
    noise = {'wm', 'winmart', 'minimart', 'hni', 'hpg', 'nan', 'pto', 'hyn', 'sla', 'ybi', 'bnh', 
             'tinh', 'thanh', 'pho', 'huyen', 'quan', 'xa', 'phuong', 'thon', 'duong', 'so', 'khu', 'cum', 'to', 'dan'}
    words = re.findall(r'\b\w{2,}\b', text) # Only words with 2+ chars
    return {w for w in words if w not in noise}

def ngram_similarity(s1, s2, n=3):
    """Calculate similarity based on character n-grams"""
    if not s1 or not s2: return 0
    def get_ngrams(s):
        return {s[i:i+n] for i in range(len(s)-n+1)}
    
    g1 = get_ngrams(s1)
    g2 = get_ngrams(s2)
    if not g1 or not g2: return 0
    return len(g1.intersection(g2)) / max(len(g1), len(g2))

def score_match(name1, addr1, name2, ma_mb2):
    """
    Combined scoring:
    1. Core keyword overlap (High weight)
    2. N-gram similarity (Medium weight)
    3. Code check (If ma_mb2 is mentioned in name1/addr1)
    """
    score = 0
    
    n1 = remove_accents(name1).lower()
    a1 = remove_accents(addr1).lower()
    n2 = remove_accents(name2).lower()
    m2 = remove_accents(ma_mb2).lower()
    
    # 1. Code check (Instant match if appraisal code is in contract name/addr)
    if m2 and (m2 in n1 or m2 in a1):
        return 1.0
    
    # 2. Keyword Overlap
    kw1 = get_core_keywords(name1 + " " + addr1)
    kw2 = get_core_keywords(name2)
    if kw1 and kw2:
        overlap = kw1.intersection(kw2)
        score += (len(overlap) / len(kw2)) * 0.6 # Focus on finding File 2 keywords in File 1
    
    # 3. Character N-gram Similarity
    score += ngram_similarity(normalize_simple(name1 + addr1), normalize_simple(name2)) * 0.4
    
    return min(score, 1.0)

def normalize_simple(text):
    text = remove_accents(text).lower()
    return "".join(re.findall(r'[a-z0-9]', text))

def main():
    print("BAT DAU: So khop du lieu nang cao V2...")
    
    try:
        wb2 = openpyxl.load_workbook(FILE_APPRAISAL, data_only=True, read_only=True)
        ws2 = wb2["01. Thẩm định giá & HSPL"]
        app_list = []
        for row in ws2.iter_rows(min_row=4, values_only=True):
            if not row or not row[4]: continue
            app_list.append({
                "ma_mb": str(row[4]),
                "info": str(row[6] or ""),
                "details": row[12],
                "gia_thue": row[13],
                "hspl": row[14]
            })
        print(f"Loaded {len(app_list)} rows from Appraisal.")

        wb1 = openpyxl.load_workbook(FILE_CONTRACT, data_only=True)
        ws1 = wb1.active
        
        # New columns
        start_col = ws1.max_column + 1
        headers = ["[V2] Ma MB", "[V2] Diem (%)", "[V2] Ten MB", "[V2] Chi tiet", "[V2] Gia", "[V2] HSPL"]
        for i, h in enumerate(headers):
            ws1.cell(row=1, column=start_col + i, value=h)

        match_count = 0
        total = ws1.max_row
        for r in range(3, total + 1):
            t_ch = str(ws1.cell(row=r, column=3).value or "")
            d_ch = str(ws1.cell(row=r, column=4).value or "")
            
            best_s = 0
            best_m = None
            
            for item in app_list:
                s = score_match(t_ch, d_ch, item["info"], item["ma_mb"])
                if s > best_s:
                    best_s = s
                    best_m = item
                if s == 1.0: break # Exact code match
            
            if best_m and best_s > 0.3: # Lower threshold but better scoring
                ws1.cell(row=r, column=start_col, value=best_m["ma_mb"])
                ws1.cell(row=r, column=start_col+1, value=round(best_s * 100, 1))
                ws1.cell(row=r, column=start_col+2, value=best_m["info"])
                ws1.cell(row=r, column=start_col+3, value=best_m["details"])
                ws1.cell(row=r, column=start_col+4, value=best_m["gia_thue"])
                ws1.cell(row=r, column=start_col+5, value=best_m["hspl"])
                match_count += 1
            
            if r % 50 == 0: print(f"   ... {r}/{total} processed")

        wb1.save(OUTPUT_FILE)
        print(f"XONG! Khop duoc {match_count} dong. Luu tai V2.")
        
    except Exception as e:
        print(f"LOI: {e}")

if __name__ == "__main__":
    main()
