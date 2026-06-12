import openpyxl
import os
import unicodedata
import re
from datetime import datetime
import sys
import io
import shutil

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

# --- CONFIGURATION ---
ORIGINAL_CONTRACT = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\danh sach ky hop dong.xlsx"
ORIGINAL_APPRAISAL = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\04. Phụ lục 03_Phân tích dữ liệu mở mới.xlsx"

# Temporary copies to avoid permission issues
TEMP_CONTRACT = r"d:\Dự án AI\chay_tool\scratch\temp_contract.xlsx"
TEMP_APPRAISAL = r"d:\Dự án AI\chay_tool\scratch\temp_appraisal.xlsx"

OUTPUT_FILE = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\Ket_Qua_Gop_Du_Lieu.xlsx"

# --- UTILS ---
def remove_accents(input_str):
    if not input_str: return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    s = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    return s.replace('đ', 'd').replace('Đ', 'D')

def normalize_name(text):
    if not text: return ""
    text = remove_accents(text).lower()
    terms_to_remove = [
        r'wm\+', r'winmart\+', r'winmart', r'minimart', r'hni', r'hpg', r'nan', r'pto', r'hyn', r'sla', r'ybi', r'bnh',
        r't\.', r'tp\.', r'h\.', r'x\.', r'p\.', r'thon', r'xa', r'huyen', r'tinh', r'phuong', r'quan', r'thanh pho',
        r'duong', r'so', r'khu', r'cum', r'to', r'dan', r'pho'
    ]
    pattern = r'\b(' + '|'.join(terms_to_remove) + r')\b'
    text = re.sub(pattern, '', text)
    text = re.sub(r'[^a-z0-9\s]', ' ', text)
    return " ".join(text.split())

def calculate_similarity(s1, s2):
    if not s1 or not s2: return 0
    words1 = set(s1.split())
    words2 = set(s2.split())
    if not words1 or not words2: return 0
    intersection = words1.intersection(words2)
    return len(intersection) / max(len(words1), len(words2))

# --- MAIN ---
def main():
    print("BAT DAU: Qua trinh gop du lieu...")
    
    # 0. Create temporary copies
    print("DANG SAO CHEP: Tao ban sao tam thoi de tranh loi file dang mo...")
    try:
        shutil.copy2(ORIGINAL_CONTRACT, TEMP_CONTRACT)
        shutil.copy2(ORIGINAL_APPRAISAL, TEMP_APPRAISAL)
    except Exception as e:
        print(f"LOI: Khong the sao chep file: {e}")
        # If copy fails, maybe the file doesn't exist or is locked too tightly
        if not os.path.exists(ORIGINAL_CONTRACT):
            print(f"File khong ton tai: {ORIGINAL_CONTRACT}")
            return

    # 1. Load Appraisal Data
    print("DANG DOC: File Tham dinh...")
    try:
        wb2 = openpyxl.load_workbook(TEMP_APPRAISAL, data_only=True, read_only=True)
        ws2 = wb2["01. Thẩm định giá & HSPL"]
        
        appraisal_list = []
        for row in ws2.iter_rows(min_row=4, values_only=True):
            if not row or row[4] is None: continue
            
            info_text = str(row[6] or "")
            norm_info = normalize_name(info_text)
            
            appraisal_list.append({
                "ma_mb": row[4],
                "info": info_text,
                "norm_info": norm_info,
                "details": row[12],
                "gia_thue": row[13],
                "hspl": row[14]
            })
        print(f"DONE: Da tai {len(appraisal_list)} dong tu file Tham dinh.")
    except Exception as e:
        print(f"LOI: Doc file Tham dinh: {e}")
        return

    # 2. Process Contract Data
    print("DANG XU LY: File Hop dong va so khop...")
    try:
        wb1 = openpyxl.load_workbook(TEMP_CONTRACT, data_only=True)
        ws1 = wb1.active
        
        start_col = ws1.max_column + 1
        headers = ["[Match] Ma MB", "[Match] Do chinh xac (%)", "[Match] Ten MB Tham dinh", 
                   "[Match] Chi tiet Tham dinh", "[Match] Gia thue", "[Match] HSPL"]
        
        for i, h in enumerate(headers):
            ws1.cell(row=1, column=start_col + i, value=h)

        match_count = 0
        total_rows = ws1.max_row
        for r_idx in range(3, total_rows + 1):
            ten_ch = str(ws1.cell(row=r_idx, column=3).value or "")
            dia_chi = str(ws1.cell(row=r_idx, column=4).value or "")
            
            full_name = f"{ten_ch} {dia_chi}"
            norm_full = normalize_name(full_name)
            
            if not norm_full: continue
            
            best_match = None
            best_score = 0
            
            for item in appraisal_list:
                score = calculate_similarity(norm_full, item["norm_info"])
                if score > best_score:
                    best_score = score
                    best_match = item
            
            if best_match and best_score > 0.2:
                ws1.cell(row=r_idx, column=start_col, value=best_match["ma_mb"])
                ws1.cell(row=r_idx, column=start_col+1, value=round(best_score * 100, 1))
                ws1.cell(row=r_idx, column=start_col+2, value=best_match["info"])
                ws1.cell(row=r_idx, column=start_col+3, value=best_match["details"])
                ws1.cell(row=r_idx, column=start_col+4, value=best_match["gia_thue"])
                ws1.cell(row=r_idx, column=start_col+5, value=best_match["hspl"])
                match_count += 1
                
            if r_idx % 50 == 0:
                print(f"   ...Da xu ly {r_idx}/{total_rows} dong")

        print(f"LUU: Dang luu ket qua vao {OUTPUT_FILE}...")
        wb1.save(OUTPUT_FILE)
        print(f"XONG: Hoan thanh! Da khop duoc {match_count} dong.")
        
    except Exception as e:
        print(f"LOI: Xu ly file Hop dong: {e}")

if __name__ == "__main__":
    main()
