import openpyxl
import os

output_path = r"D:\chạy tool\chạy tool\LICH_SU_PHAT_TRIEN_TOOL.xlsx"

if os.path.exists(output_path):
    try:
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Clear current highlight
        for r in range(ws.max_row, ws.max_row+1):
            for c in range(1,7):
                ws.cell(row=r, column=c).fill = PatternFill(fill_type=None)
                
        new_row = [
            11, "10:15 AM", "Cải tiến Regex Thông minh bóc tách Tên Tỉnh (Fix lỗi bôi vàng)",
            "Chẩn đoán pháp y nội dung email, xác nhận cấu trúc mới: Tỉnh nằm sát ngoặc đơn ở tiêu đề.",
            "Tạo Snapshot bảo hiểm. Bổ sung Regex bóc cụm từ ngăn cách sát mã MB. Nâng tỉ lệ quét Tỉnh lên tối đa.",
            "HOÀN THÀNH & ĐÃ ĐỒNG BỘ"
        ]
        ws.append(new_row)
        idx = ws.max_row
        for col in range(1, 7):
            c = ws.cell(row=idx, column=col)
            c.alignment = left_align
            c.border = thin_border
            c.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
            
        wb.save(output_path)
    except: pass
