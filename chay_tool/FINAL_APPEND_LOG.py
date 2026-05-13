import openpyxl
import os

output_path = r"D:\chạy tool\chạy tool\LICH_SU_PHAT_TRIEN_TOOL.xlsx"

if os.path.exists(output_path):
    try:
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Clear color from previous lines
        for r in range(ws.max_row-1, ws.max_row+1):
            for c in range(1,7):
                ws.cell(row=r, column=c).fill = PatternFill(fill_type=None)

        records = [
            [8, "09:15 AM", "Lệnh Khôi Phục nhầm về 06:59 AM & Hoàn tác về v2.5.15", 
             "Mất ổn định tạm thời do lệch cấu trúc Column mapping và data template cũ.", 
             "Nhận diện xung đột logic và template excel.", "Chờ sửa lỗi"],
            [9, "09:25 AM", "BÁO CÁO: Dữ liệu Excel nhân bản & lệch cột nửa dưới", 
             "Phát hiện sai lệch giữa Code mới (xóa cột) và File Mẫu cũ (vẫn còn Header).", 
             "Thực hiện script Tự động XÓA RÁC dòng 83+ và REVERT Code về bản chuẩn 100%.", "Đã dọn dẹp sạch sẽ"],
            [10, "HIỆN TẠI (09:55 AM)", "THỰC THI CHÍNH THỨC: Cắt bỏ triệt để cột Trưởng Nhóm (Ủy quyền tay 100%)", 
             "Hoàn tất Combo 3 trong 1: Sửa Logic Core + Sửa UI Excel Mapping + Trảm vật lý Cột M trên File Mẫu.", 
             "Tạo Snapshot siêu dự phòng. Xóa code thừa. Xóa Column 13 template. Khớp cấu trúc Tuyệt Đối.", "HOÀN THÀNH TỐI ƯU MỚI"]
        ]
        
        for record in records:
            ws.append(record)
            idx = ws.max_row
            for col in range(1, 7):
                c = ws.cell(row=idx, column=col)
                c.alignment = left_align
                c.border = thin_border
                # Highlight only the very last current record
                if record[0] == 10:
                    c.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
                    
        wb.save(output_path)
    except:
        pass
