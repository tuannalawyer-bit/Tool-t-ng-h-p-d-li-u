"""
merge_final_v2.py
-----------------
Ghép 2 file Excel theo địa chỉ (Phiên bản cải tiến - Chống false positives).
Khắc phục lỗi:
1. Lỗi ngoặc đơn chứa số: (1), (2) bị nhận nhầm làm địa chỉ đầu vào.
2. Lỗi token_set_ratio cho điểm 100% sai do các từ hành chính ("Xã", "Phường", "Thành phố") bị trùng lặp.
-> Giải pháp: Loại bỏ toàn bộ từ hành chính trước khi so khớp để chỉ so sánh "Tên Lõi".
"""

import pandas as pd
import sys, io, re
from rapidfuzz import fuzz, process
from unidecode import unidecode

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATH1 = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\danh sach ky hop dong.xlsx"
PATH2 = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\form tool.xlsx"
OUTPUT = r"D:\OneDrive - WIN\Desktop\công việc thực hiện\Tháng 5\Chuyên đề xây dựng mở mới\merged_result_v2_fixed.xlsx"

TINH_THRESHOLD = 80
ADDR_THRESHOLD = 75
EXACT_THRESHOLD = 90

ADMIN_WORDS = [
    r'\bthanh pho\b', r'\btp\b', r'\btinh\b', r'\bquan\b', r'\bhuyen\b', r'\bthi xa\b', r'\bthi tran\b', r'\btt\b',
    r'\bphuong\b', r'\bxa\b', r'\bthon\b', r'\bxom\b', r'\bto dan pho\b', r'\btdp\b', r'\bto\b', r'\bkhu vuc\b',
    r'\bkhu\b', r'\bap\b', r'\bduong\b', r'\bpho\b', r'\bngo\b', r'\bhem\b', r'\bso\b', r'\btoa nha\b', r'\bcho\b',
    r'\bkdt\b', r'\bkhu do thi\b', r'\bd\b', r'\bp\b', r'\bq\b', r'\bt\b', r'\bh\b', r'\bx\b', r'\btb\b', r'\btbd\b', r'\btd\b'
]

def normalize_tinh(s):
    if pd.isna(s): return ""
    s = str(s).strip()
    s = re.sub(r'^(TP\.|T\.|Tỉnh|Thành phố|Thành Phố|TPHCM)\s*', '', s, flags=re.IGNORECASE)
    s = s.split('/')[0].strip()
    return unidecode(s).lower().strip()

def normalize_core_addr(s):
    """Normalize & remove administrative words"""
    if pd.isna(s): return ""
    s = s.lower()
    s = unidecode(s)
    s = re.sub(r'[\W_]+', ' ', s)
    
    for word in ADMIN_WORDS:
        s = re.sub(word, ' ', s)
    
    return " ".join(s.split())

def extract_outer_addr(s):
    if pd.isna(s): return ""
    s = re.sub(r'\([^)]*\)', '', str(s)) # Remove anything in brackets
    return normalize_core_addr(s)

def extract_inner_addr(s):
    if pd.isna(s): return ""
    matches = re.findall(r'\(([^)]+)\)', str(s))
    for inner in matches:
        if len(inner) < 5: # Ignore "(1)", "(2)"
            continue
        inner = re.sub(r'^(?:Địa chỉ (?:đầu vào|ban đầu)|Đầu vào)[:\s]*', '', inner, flags=re.IGNORECASE)
        return normalize_core_addr(inner)
    return ""

print("Đang tải dữ liệu...")
df1 = pd.read_excel(PATH1)
df2 = pd.read_excel(PATH2, header=2)

df1 = df1.dropna(subset=['Địa chỉ']).reset_index(drop=True)
df2 = df2.dropna(subset=['Thông tin MB đang thuê']).reset_index(drop=True)

print("Đang chuẩn hóa dữ liệu (Tách từ khóa lõi)...")
tinh1_n = df1['Tỉnh/TP/Quận'].apply(normalize_tinh)
tinh2_n = df2['Tỉnh'].apply(normalize_tinh)
addr1_n = df1['Địa chỉ'].apply(normalize_core_addr)
addr2_outer_n = df2['Thông tin MB đang thuê'].apply(extract_outer_addr)
addr2_inner_n = df2['Thông tin MB đang thuê'].apply(extract_inner_addr)

tinh2_unique = tinh2_n.unique().tolist()
addr2_outer_list = addr2_outer_n.tolist()
addr2_inner_list = addr2_inner_n.tolist()

match_df2_idx   = []
match_scores    = []
match_statuses  = []
match_methods   = []

print("Đang so khớp...")
for idx1 in range(len(df1)):
    t1 = tinh1_n.iloc[idx1]
    a1 = addr1_n.iloc[idx1]

    best_df2_idx = None
    best_score = 0
    method = "NO_MATCH"

    if len(a1.strip()) < 3: # Skip empty or too short core addresses
        match_df2_idx.append(None)
        match_scores.append(0)
        match_statuses.append("NO_MATCH")
        match_methods.append("NO_MATCH")
        continue

    tinh_match = process.extractOne(t1, tinh2_unique, scorer=fuzz.ratio)

    if tinh_match and tinh_match[1] >= TINH_THRESHOLD:
        matched_tinh_val = tinh_match[0]
        sub_idx = tinh2_n[tinh2_n == matched_tinh_val].index.tolist()

        sub_outer = [addr2_outer_list[j] for j in sub_idx]
        r_out = process.extractOne(a1, sub_outer, scorer=fuzz.token_set_ratio)

        sub_inner_pairs = [(addr2_inner_list[j], j) for j in sub_idx if addr2_inner_list[j]]
        r_inn = None
        if sub_inner_pairs:
            r_inn = process.extractOne(a1, [v for v, _ in sub_inner_pairs], scorer=fuzz.token_set_ratio)

        score_out = r_out[1] if r_out else 0
        score_inn = r_inn[1] if r_inn else 0

        # Safety check: token_set_ratio can be 100 if one is subset.
        # Use a secondary scorer to penalize if one string is vastly different
        if score_out >= score_inn and score_out >= ADDR_THRESHOLD:
            # secondary check with token_sort_ratio to avoid subset traps
            sec_score = fuzz.token_sort_ratio(a1, sub_outer[r_out[2]])
            final_score = (score_out + sec_score) / 2 # Blended score is much safer!
            if final_score >= ADDR_THRESHOLD:
                best_df2_idx = sub_idx[r_out[2]]
                best_score = final_score
                method = f"TINH+OUTER"

        if score_inn > score_out and score_inn >= ADDR_THRESHOLD:
            sec_score = fuzz.token_sort_ratio(a1, sub_inner_pairs[r_inn[2]][0])
            final_score = (score_inn + sec_score) / 2
            if final_score >= ADDR_THRESHOLD and final_score > best_score:
                best_df2_idx = sub_inner_pairs[r_inn[2]][1]
                best_score = final_score
                method = f"TINH+INNER"

    # Fallback
    if best_df2_idx is None:
        r_global = process.extractOne(a1, addr2_outer_list, scorer=fuzz.token_set_ratio)
        if r_global and r_global[1] >= ADDR_THRESHOLD:
            sec_score = fuzz.token_sort_ratio(a1, addr2_outer_list[r_global[2]])
            final_score = (r_global[1] + sec_score) / 2
            if final_score >= ADDR_THRESHOLD:
                best_df2_idx = r_global[2]
                best_score = final_score
                method = "GLOBAL_FALLBACK"

    if best_df2_idx is None:
        status = "NO_MATCH"
        best_score = 0
    elif best_score >= EXACT_THRESHOLD:
        status = "EXACT"
    else:
        status = "FUZZY"

    match_df2_idx.append(best_df2_idx)
    match_scores.append(round(best_score, 1))
    match_statuses.append(status)
    match_methods.append(method)

n_exact = match_statuses.count("EXACT")
n_fuzzy = match_statuses.count("FUZZY")
n_none  = match_statuses.count("NO_MATCH")
n_total = len(df1)

print(f"\n📊 Kết quả so khớp:")
print(f"   EXACT  (score ≥{EXACT_THRESHOLD}%): {n_exact:3d} dòng ({n_exact/n_total*100:.1f}%)")
print(f"   FUZZY  ({ADDR_THRESHOLD}-{EXACT_THRESHOLD}%):      {n_fuzzy:3d} dòng ({n_fuzzy/n_total*100:.1f}%)")
print(f"   NO_MATCH (<{ADDR_THRESHOLD}%):      {n_none:3d} dòng ({n_none/n_total*100:.1f}%)")
print(f"   TỔNG KHỚP:          {n_exact+n_fuzzy:3d} dòng ({(n_exact+n_fuzzy)/n_total*100:.1f}%)")

df1['_match_idx']    = match_df2_idx
df1['_match_score']  = match_scores
df1['_match_status'] = match_statuses
df1['_match_method'] = match_methods

df2_renamed = df2.copy()
df2_renamed.columns = [f"TĐG_{c}" if c not in ['STT'] else f"TĐG_{c}" for c in df2.columns]

df2_matched_rows = []
for row_idx in match_df2_idx:
    if row_idx is not None:
        df2_matched_rows.append(df2_renamed.iloc[row_idx].to_dict())
    else:
        df2_matched_rows.append({c: None for c in df2_renamed.columns})

df2_matched = pd.DataFrame(df2_matched_rows)
df2_matched.index = df1.index

merged = pd.concat([df1, df2_matched], axis=1)
merged = merged.rename(columns={
    '_match_score':  'MATCH_SCORE (%)',
    '_match_status': 'MATCH_STATUS',
    '_match_method': 'MATCH_METHOD',
})
merged = merged.drop(columns=['_match_idx'], errors='ignore')

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
    merged.to_excel(writer, index=False, sheet_name='Kết quả ghép')
    ws = writer.sheets['Kết quả ghép']
    from openpyxl.styles import PatternFill
    fill_exact = PatternFill("solid", fgColor="C6EFCE")
    fill_fuzzy = PatternFill("solid", fgColor="FFEB9C")
    fill_none  = PatternFill("solid", fgColor="FFC7CE")
    status_col = merged.columns.get_loc('MATCH_STATUS') + 1
    for row_i, status in enumerate(match_statuses, start=2):
        fill = fill_exact if status == "EXACT" else (fill_fuzzy if status == "FUZZY" else fill_none)
        ws.cell(row=row_i, column=status_col).fill = fill

print(f"\n✅ File kết quả lưu tại:\n   {OUTPUT}")
