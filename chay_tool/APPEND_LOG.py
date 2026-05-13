import openpyxl
import os

output_path = r"D:\chạy tool\chạy tool\LICH_SU_PHAT_TRIEN_TOOL.xlsx"

if os.path.exists(output_path):
    try:
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Un-highlight row 6
        for c in range(1,7):
            ws.cell(row=6, column=c).fill = PatternFill(fill_type=None)
            
        new_row = [
            6, "HIỆN TẠI (08:55 AM)", "Yêu cầu loại bỏ hoàn toàn cột Trưởng Nhóm Khỏi Hệ Thống.",
            "Đã phẫu thuật thành công Logic code, AI và Tự động vá file Template Excel Mẫu.",
            "Tạo Snapshot Snapshot, Xóa biến tn_tkmb, Dồn dịch cột Excel, Tự động xóa cột M trong file mẫu.",
            "Phiên bản Tối Ưu Mới"
        ]
        ws.append(new_row)
        
        idx = ws.max_row
        for col in range(1, 7):
            c = ws.cell(row=idx, column=col)
            c.alignment = left_align
            c.border = thin_border
            c.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
            
        wb.save(output_path)
    except: pass
